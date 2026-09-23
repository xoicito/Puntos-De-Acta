import base64
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from config import (
    ACTA_BOARD_ID,
    ACTA_OUTPUT_DIR,
    ACTA_XLSX_COLUMN_ID,
    FIRMA_BOARD_ID,
    FIRMA_PA_EDITABLE_COLUMN_ID,
    GERENTES_BOARD_ID,
    GERENTE_EMAIL_COLUMN_ID,
    GERENTE_EMAIL_LINK_COLUMN_ID,
    GERENTE_FIRMA_ESTADO_COLUMN_ID,
    GERENTE_FIRMA_ESTADO_FIRMADO,
    GERENTE_FIRMA_ESTADO_PENDIENTE,
    GERENTE_FIRMA_PLACEHOLDER,
    GERENTE_LINK_BASE_URL,
    GERENTE_FIRMA_LINK_COLUMN_ID,
    GERENTE_NOMBRE_COLUMN_ID,
    PMO_BOARD_ID,
    PMO_EMAIL_APROBACION_COLUMN_ID,
    PMO_EMAIL_COLUMN_ID,
    PMO_NOMBRE_COLUMN_ID,
    TEST_MODE_SKIP_NOTIFICATIONS,
)
from gerente_link import InvalidLinkError, generate_signing_link, verify_token
from utils.acta_builder import item_data
from utils.excel_writer import insert_signature
from utils.monday_client import (
    change_status,
    create_item,
    create_update,
    find_item_by_name,
    get_file_public_url,
    get_item,
    download_file,
    update_text_column,
    upload_file,
)


def _column_text(item, column_id):
    if not column_id:
        return ""

    for c in item.get("column_values", []):
        if c["id"] == column_id:
            return (c.get("text") or "").strip()

    return ""


def start_gerente_signing(item_id, board_id):
    """Called right after a Punto de Acta is generated: resolve the real
    Gerente de Proyecto by looking up the name the Lider picked (a fixed
    dropdown - Connect Boards columns don't work in Monday's public
    forms) against the "Gerentes" board, send him a signing link, and
    notify him inside Monday too.

    No-ops quietly (just logs) if the name column isn't configured yet
    or nobody's selected - this is meant to be safe to call unconditionally
    from generate_acta.py.
    """

    if not GERENTE_NOMBRE_COLUMN_ID or not GERENTE_EMAIL_COLUMN_ID:
        print("GERENTE_FIRMA: columnas de Gerente no configuradas, se omite")
        return

    if TEST_MODE_SKIP_NOTIFICATIONS:
        print("GERENTE_FIRMA: TEST_MODE_SKIP_NOTIFICATIONS activo, se omite enlace de firma y todo lo posterior")
        return

    item = get_item(item_id)
    selected_name = _column_text(item, GERENTE_NOMBRE_COLUMN_ID)
    name, email = find_item_by_name(GERENTES_BOARD_ID, selected_name, GERENTE_EMAIL_COLUMN_ID)

    if not email:
        print(f"GERENTE_FIRMA: item={item_id} sin Gerente de Proyecto asignado (o '{selected_name}' no encontrado), se omite")
        return

    link = generate_signing_link(item_id, board_id)

    if GERENTE_FIRMA_LINK_COLUMN_ID:
        update_text_column(item_id, board_id, GERENTE_FIRMA_LINK_COLUMN_ID, link)

    if GERENTE_EMAIL_LINK_COLUMN_ID:
        update_text_column(item_id, board_id, GERENTE_EMAIL_LINK_COLUMN_ID, email)

    # Se marca hasta el final, ya con el link y el correo escritos - la
    # automatizacion que le avisa al Gerente debe disparar sobre este
    # cambio, para no mandar el correo antes de que "Gerente Correo"
    # tenga la direccion (mismo problema que ya se arreglo para PMO).
    if GERENTE_FIRMA_ESTADO_COLUMN_ID:
        change_status(item_id, board_id, GERENTE_FIRMA_ESTADO_COLUMN_ID, GERENTE_FIRMA_ESTADO_PENDIENTE)

    try:
        create_update(
            item_id,
            f"El Punto de Acta esta listo para la firma de {name or 'Gerente de Proyecto'}. "
            f"Enlace de firma: {link}",
        )
    except Exception as e:
        print(f"GERENTE_FIRMA: no se pudo publicar el update en Monday: {e}")

    print(f"GERENTE_FIRMA: enlace generado para {name} <{email}> item={item_id}")


def _current_estado(item):
    if not GERENTE_FIRMA_ESTADO_COLUMN_ID:
        return ""

    for c in item.get("column_values", []):
        if c["id"] == GERENTE_FIRMA_ESTADO_COLUMN_ID:
            return (c.get("text") or "").strip()

    return ""


def _reject_if_already_signed(item):
    if _current_estado(item) == GERENTE_FIRMA_ESTADO_FIRMADO:
        raise LookupError("Este Punto de Acta ya fue firmado.")


def resolve_signing_context(token):
    """Validate the token and load what the signing page needs to show.

    Raises InvalidLinkError (bad/expired token) or LookupError (already
    signed) with a message safe to display to the signer.
    """

    data = verify_token(token)
    item_id = data["item_id"]
    board_id = data["board_id"]

    item = get_item(item_id)
    data_fields = item_data(item)

    _reject_if_already_signed(item)

    return {
        "item_id": item_id,
        "board_id": board_id,
        "proyecto": data_fields.get("proyecto", ""),
        "no_contrato": data_fields.get("no_contrato", ""),
        "empresa": data_fields.get("empresa", ""),
        "tipo_plantilla": (data_fields.get("tipo_plantilla") or "").strip().lower(),
    }


def get_current_document_url(token):
    """Resolve a fresh public URL for the acta as the Lider left it, so
    the signing page can offer a "download before you sign" link. Re-runs
    the token check so this can't be used to peek at a document the
    caller never had a valid link for.
    """

    data = verify_token(token)
    item = get_item(data["item_id"])

    return get_file_public_url(item, ACTA_XLSX_COLUMN_ID)


def _save_signature_image(output_directory, stem, file_storage=None, data_url=None):
    if file_storage is not None and file_storage.filename:
        suffix = Path(file_storage.filename).suffix or ".png"
        path = str(output_directory / f"{stem}_firma_gerente{suffix}")
        file_storage.save(path)
        return path

    if data_url:
        header, _, encoded = data_url.partition(",")
        if not encoded:
            raise ValueError("Firma dibujada vacia")
        image_bytes = base64.b64decode(encoded)
        path = str(output_directory / f"{stem}_firma_gerente.png")
        Path(path).write_bytes(image_bytes)
        return path

    raise ValueError("No se recibio ninguna firma")


def _send_to_procurement(item_id, item, signed_path, data_fields):
    """Per the approval policy, once both Lider and Gerente have signed (the
    'PA Inicial'), the request moves to Procurement (Arq. Melissa Alvarenga's
    board) for review. Creates a new item there and attaches the doubly-
    signed file to PA EDITABLE - this replaces the old manual upload step.

    Also carries over the PMO's email (resolved here, from the name the
    Lider picked on the *original* item - the new Procurement item has no
    such column of its own) into PMO_EMAIL_APROBACION_COLUMN_ID on the new
    item. The PMO isn't notified yet at this point - that only happens
    once Melissa signs, via a Monday automation on the Procurement board
    watching that column, same pattern as everywhere else.
    """

    if not FIRMA_BOARD_ID or not FIRMA_PA_EDITABLE_COLUMN_ID:
        print("GERENTE_FIRMA: board/columna de Procurement no configurados, se omite envio")
        return

    acta_id = data_fields.get("acta_id") or f"item-{item_id}"
    proyecto = data_fields.get("proyecto") or ""
    rubro = data_fields.get("rubro") or ""

    # El board de Aprobacion no tiene una columna "Rubro" propia - el
    # Rubro elegido en el Forms PA se refleja en el nombre del item
    # (la columna "name"), igual que Acta ID y Proyecto.
    name = " - ".join(part for part in (acta_id, proyecto, rubro) if part)

    new_item_id = create_item(FIRMA_BOARD_ID, name)

    upload_file(new_item_id, FIRMA_PA_EDITABLE_COLUMN_ID, signed_path)

    if PMO_NOMBRE_COLUMN_ID and PMO_EMAIL_COLUMN_ID and PMO_EMAIL_APROBACION_COLUMN_ID:
        pmo_selected = _column_text(item, PMO_NOMBRE_COLUMN_ID)
        pmo_name, pmo_email = find_item_by_name(PMO_BOARD_ID, pmo_selected, PMO_EMAIL_COLUMN_ID)

        if pmo_email:
            update_text_column(new_item_id, FIRMA_BOARD_ID, PMO_EMAIL_APROBACION_COLUMN_ID, pmo_email)
            print(f"PMO: correo resuelto para {pmo_name} <{pmo_email}>, copiado a item={new_item_id}")
        else:
            print(f"PMO: item={item_id} sin PMO asignado (o '{pmo_selected}' no encontrado), se omite")

    print(f"GERENTE_FIRMA: enviado a Procurement, item={new_item_id} ({name})")


def apply_gerente_signature(token, file_storage=None, data_url=None, audit=None):
    """Verify the token again, insert the signature, upload the result, and
    mark the item as signed. Meant to run on the POST of the signing page -
    never trust the GET's validation alone.
    """

    data = verify_token(token)
    item_id = data["item_id"]
    board_id = data["board_id"]

    item = get_item(item_id)
    _reject_if_already_signed(item)

    editable_url = get_file_public_url(item, ACTA_XLSX_COLUMN_ID)

    if not editable_url:
        raise ValueError("El item no tiene un acta generada todavia")

    output_directory = Path(ACTA_OUTPUT_DIR)
    output_directory.mkdir(parents=True, exist_ok=True)

    stem = f"item_{item_id}"

    suffix = Path(editable_url.split("?")[0]).suffix or ".xlsx"
    source_path = str(output_directory / f"{stem}_actual{suffix}")
    download_file(editable_url, source_path)

    signature_path = _save_signature_image(
        output_directory, stem, file_storage=file_storage, data_url=data_url
    )

    wb = load_workbook(source_path)
    ws = wb["C-9-12"] if "C-9-12" in wb.sheetnames else wb.active

    signed_ok = insert_signature(ws, signature_path, GERENTE_FIRMA_PLACEHOLDER)

    if not signed_ok:
        raise RuntimeError("No se pudo insertar la firma en el documento")

    signed_path = str(output_directory / f"{stem}_firmado.xlsx")
    wb.save(signed_path)

    upload_file(item_id, ACTA_XLSX_COLUMN_ID, signed_path)

    try:
        _send_to_procurement(item_id, item, signed_path, item_data(item))
    except Exception as e:
        print(f"GERENTE_FIRMA: no se pudo enviar a Procurement: {e}")

    if GERENTE_FIRMA_ESTADO_COLUMN_ID:
        change_status(item_id, board_id, GERENTE_FIRMA_ESTADO_COLUMN_ID, GERENTE_FIRMA_ESTADO_FIRMADO)

    audit = audit or {}
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    # Written into Monday's own activity feed rather than just printed, so
    # this survives Render's log rotation - it's the actual audit record.
    try:
        create_update(
            item_id,
            f"Documento firmado por el Gerente de Proyecto el {timestamp}. "
            f"IP: {audit.get('ip', 'desconocida')} - "
            f"Dispositivo: {audit.get('user_agent', 'desconocido')}",
        )
    except Exception as e:
        print(f"GERENTE_FIRMA: no se pudo publicar el registro de auditoria: {e}")

    print(f"GERENTE_FIRMA: item={item_id} firmado - {audit}")

    data_fields = item_data(item)

    return {
        "item_id": item_id,
        "tipo_plantilla": (data_fields.get("tipo_plantilla") or "").strip().lower(),
    }
