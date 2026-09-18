from openpyxl import load_workbook


def parse_cotizacion_upload(path):
    """Read a filled-in copy of PLANTILLA_ALCANCE_COTIZACION.xlsx and return
    a list of row dicts shaped for write_cotizacion_rows() (descripcion,
    unidad, cantidad, precio, observaciones).

    Any row with a non-empty Descripcion (column B) is included, in
    whatever order it appears - blank rows in between are just skipped, so
    it doesn't matter if the Lider leaves gaps or adds rows past the
    template's original 25.
    """

    wb = load_workbook(path, data_only=True)
    ws = wb.active

    rows = []

    for row in ws.iter_rows(min_row=5):
        descripcion = row[1].value  # column B

        if not descripcion or not str(descripcion).strip():
            continue

        rows.append({
            "descripcion": str(descripcion).strip(),
            "unidad": _text(row[2].value),
            "cantidad": _number(row[3].value),
            "precio": _number(row[4].value),
            "observaciones": _text(row[6].value),
        })

    return rows


def _text(value):
    return str(value).strip() if value is not None else ""


def _number(value):
    if value is None or value == "":
        return 0

    try:
        return float(value)
    except (TypeError, ValueError):
        return 0
