from pathlib import Path

from openpyxl import load_workbook

from config import (
    ACTA_OUTPUT_DIR,
    FIRMA_BOARD_ID,
    FIRMA_DOCUMENTO_LISTO_COLUMN_ID,
    FIRMA_DOCUMENTO_LISTO_LABEL,
    FIRMA_NOTIFICAR_NOMBRES,
    FIRMA_PA_EDITABLE_COLUMN_ID,
    FIRMA_PA_FIRMADO_COLUMN_ID,
    FIRMA_PLACEHOLDER,
    MELISSA_SIGNATURE_PATH,
    TIPO_CONTRATO_MENSAJE_PLACEHOLDER,
)
from utils.monday_client import (
    change_status,
    create_update,
    download_file,
    get_file_public_url,
    get_item,
    upload_file,
)
from utils.excel_writer import insert_signature, insert_text_placeholder
from utils.pdf_convert import convert_to_pdf


def sign_document(item_id, tipo_contrato_mensaje=""):
    item = get_item(item_id)

    editable_url = get_file_public_url(item, FIRMA_PA_EDITABLE_COLUMN_ID)

    if not editable_url:
        raise ValueError(
            f"El item {item_id} no tiene archivo en la columna PA EDITABLE"
        )

    base = Path(__file__).resolve().parent

    output_directory = Path(ACTA_OUTPUT_DIR)
    output_directory.mkdir(parents=True, exist_ok=True)

    suffix = Path(editable_url.split("?")[0]).suffix or ".xlsx"
    source_path = str(output_directory / f"PA_EDITABLE_{item_id}{suffix}")

    download_file(editable_url, source_path)

    wb = load_workbook(source_path)
    ws = wb["C-9-12"] if "C-9-12" in wb.sheetnames else wb.active

    signed_ok = insert_signature(ws, str(base / MELISSA_SIGNATURE_PATH), FIRMA_PLACEHOLDER)

    if not signed_ok:
        raise RuntimeError(
            f"No se pudo insertar la firma en el item {item_id}; "
            "no se sube el archivo a PA FIRMADO PRC"
        )

    if tipo_contrato_mensaje:
        insert_text_placeholder(ws, tipo_contrato_mensaje, TIPO_CONTRATO_MENSAJE_PLACEHOLDER)

    signed_path = str(output_directory / f"PA_FIRMADO_{item_id}.xlsx")
    wb.save(signed_path)

    # Solo el PDF final va a PA FIRMADO PRC - el Excel ya no se sube ahi.
    # No esta envuelto en try/except a proposito: si la conversion falla,
    # sign_document() debe fallar tambien (nada que subir), en vez de
    # dejar la columna vacia y de todos modos marcar "Enviado a PMO"
    # como Listo mas abajo.
    pdf_path = convert_to_pdf(signed_path, output_directory)
    upload_file(item_id, FIRMA_PA_FIRMADO_COLUMN_ID, pdf_path)

    # Se marca hasta el final, ya con el PDF subido - la automatizacion
    # que le manda el correo a PMO con el adjunto debe disparar sobre
    # este cambio, no sobre ESTADO DE APROBACION (ver config.py), para
    # no mandar el correo antes de que el archivo este listo.
    if FIRMA_DOCUMENTO_LISTO_COLUMN_ID:
        change_status(item_id, FIRMA_BOARD_ID, FIRMA_DOCUMENTO_LISTO_COLUMN_ID, FIRMA_DOCUMENTO_LISTO_LABEL)

    if FIRMA_NOTIFICAR_NOMBRES:
        try:
            nombres = ", ".join(FIRMA_NOTIFICAR_NOMBRES)
            create_update(
                item_id,
                f"Arq. Melissa Alvarenga firmo el Punto de Acta. "
                f"Notificando a: {nombres}.",
            )
        except Exception as e:
            print(f"FIRMA_MELISSA: no se pudo publicar la notificacion: {e}")

    return {"signed": signed_path}


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("item_id")
    args = parser.parse_args()

    print(sign_document(args.item_id))
